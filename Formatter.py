from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries 
from openpyxl.worksheet.Table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from typing import Optional, Literal   

from ValidatorTool_Congfig.ColorMaps import (
    ERROR_COLORS, HEADER_COLORS, 
    EXTRA_ERRORS_STYLE, METADATA_STYLE
)
from ValidatorTool_Congfig.MetadataMaps import combine_error_maps 

class Formatter: 
    """
    Provides specialized formatting to Excel Workbook, 
    used by *ProcessMetadata*.

    Attributes: 
        all_errors_map (dict[int,list[str]]): 
            All error maps in one dictionary 
    """

all_errors_map = combine_error_maps() 

def format(self) -> None: 
    """
    Applies color and table formatting, and created new tab for Legend/Extra Errors. 
    """
    # ---Fill Mapping Worksheet specification variable---
    ws_full_map = self.workbook[self.full_map_sheet]
    # ---Basic Styling---
    self.add_all_borders(
        ws_full_map, 
        line_color="D9D9D9"
    )

    #---Header Coloring Convention---
    for header in ws_full_map[1]: 
        cell = header.value 
        # ---"sql_path" as a safety measure if failed to be removed--
        if cell in ("sql_path", "ErrorCodes"): 
            continue
        header.fill = PatternFill(
            fill_type="solid", 
            start_color=HEADER_COLORS.get[cell], 
            end_color=HEADER_COLORS.get[cell]
    )
        header.font = Font(bold=True, color="000000")
        header.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
    # ---ErrorCodes fill colors---
    header_letter_map = self._get_header_letter_map(ws_full_map, "letters")
    if "f" in self.colorblind: 
        for error_type in ERROR_COLORS: 
            ERROR_COLORS[error_type]["color"] = "A6A6A6"
    fills = {
        error_type: PatternFill(
            start_code=hex_code["color"],
            end_code=hex_code["color"],
            fill_type="solid"
        ) 
    for error_type, hex_code in ERROR_COLORS.items()
    }
    # ---Parse ErrorCodes--- 
    header_indexes_map = self._get_headers(ws_full_map, "indexes")
    error_codes_col_inx = header_indexes_map["ErrorCodes"]
    for row_inx in range(2, ws_full_map.max_row + 1): 
        cell_value = ws_full_map.cell(row=row_inx, column=error_codes_col_inx).value
        if not cell_value: 
            continue
        # ---Separate ErrorCodes -- 
        error_codes = [code.strip() for code in cell_value.split(",") if code.strip()]
        for code_str in error_codes: 
            # --- Convert & Compare each code to all_errors_map dict ---
            try: 
                code_key = int(code_str)
            except ValueError: 
                code_key = code_str
            # --- Diversion if error code does not exist --- 
            if code_key not in self.all_errors_map: 
                continue
            # --- 1st digit determine error color--- 
            prefix = fills.get(prefix)
            if not fill: 
                continue 
            # --- Find specific column for each code --- 
            target_columns = self.all_errors_map[code_key]
            # --- Handling Key: Valye differences in all_errors_map ---
            if isinstance(target_columns, str):
                columns_to_highlight = [target_columns]
            else: 
                columns_to_highlight = list(target_columns)
            for col_name in columns_to_highlight:
                clean_col_name = str(col_name)
                if clean_col_name in header_letter_map: 
                    col_letter = header_letter_map[clean_col_name]
                    ws_full_map[f"{col_letter}{row_inx}"]
                    cell.fill = fill 
            # -- Create Final Table --- 
            ws_full_map.delete_cols(error_codes_col_inx)
            self.create_table(
                "ValidationResults", 
                ws_full_map,
                col_range="A:AI", 
                style=METADATA_STYLE
            )
            self.autofit_columns(ws_full_map, exceptions=("element definitions"))
            # --- Create Extra Errors Sheet -- 
            data_issues_sheet_name = "Extra Errors"
            if data_issues_sheet_name in self.workbook.sheetnames: 
                del self.workbook[data_issues_sheet_name]
            ws_extra_errors = self.workbook.create_sheet(data_issues_sheet_name)
            # -- Create legend at the top of sheet --- 
            data_issues_start_row = 1 
            if not "f" in self.colorblind: 
                for row, error_info in enumerate(ERROR_COLORS.values(), start=1): 
                    # --- Description --- 
                    ws_extra_errors[f"A{row}"] = error_info["description"]
                    # --- Corresponding Color --- 
                    ws_extra_errors[f"B{row}"].fill = PatternFill(
                        fill_type="solid",
                        start_color=error_info["color"],
                        end_color=error_info["color"],
                    )
                    # --- Add row leniency --- 
                    ws_extra_errors.append([])
                    data_issues_start_row = 5 
                # --- Format new table if needed --- 
                if self.data_issues: 
                    ws_extra_errors.append(["SQL Path", "Notes"]) 
                    for sql_path, issue_notes in self.data_issues.items(): 
                        for note in issue_notes:
                            ws_extra_errors.append([sql_path, note])
                    self.create_table(
                        "ExtraErrors", 
                        ws_extra_errors, 
                        col_range="A:B", 
                        start_row=data_issues_start_row,
                        style=EXTRA_ERRORS_STYLE
                    )
                self.autofit_columns(ws_extra_errors)

    def _get_headers(
        self, 
        sheet: Worksheet,
        header_type: Literal["letters", "indexes"] 
        ) -> dict[str, str | int]: 
        """**Private Method**.
        Gathers headers and corresponding *letter* or *index*. 

        Args: 
            sheet (Worksheet): 
                sheet required to pull heades 
            header_type ("letters", "indexes"):
                Determines whether corresponding 
                *letter* or *index* values are returned. 

        Returns: 
            dict: 
                Mapped headers with specificed *key*:*value* markers. 
        """
        header_map = {}
        for col_inx, cell in range(1, sheet.max_column + 1): 
           header = sheet.cell(row=1, column=col_inx).value
           if not header: 
                continue
           if header_type == "letters":
                header_map[header] = get_column_letter(col_inx)
           else: 
                header_map[header] = col_inx
        return header_map 

    def create_table(
        self, 
        table_name: str, 
        sheet: Worksheet, 
        col_range: Optional[str] = None, 
        start_row: int = 1,
        style: str = "TableStyleLight1",
    ) -> None: 
        """
        Create and format an Excel Worksheet into a table. 

        Args: 
            table_name (str): 
                Given name for table being created
            sheet (Worksheet):
                Sheet required to perform given required actions
            col_range (opt | str): 
                Specificed column range of table (*ex. "A:AA"*)
            start_row (int): 
                Row where table starts, (*default is row 1*)
            style (str):
                Excel given table style, (*default is "TableStyleLight1"*)
        """
        # --- Parse the col_range if applicable ---
        if col_range:
            start_col, end_col = col_range.split(":")
            range = f"{start_col}{start_row}:{end_col}{sheet.max_row}"
        else:
            range = sheet.dimensions
        # Dimensions setup -- 
        table = Table(displayName=table_name, ref=range)
        # --- Table styling --- 
        if style == "":
            style = "TableStyleLight1"
        style = TableStyleInfo(
            name=style,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        table.tableStyleInfo = style
        sheet.add_table(table)

    def add_all_borders(
        self, 
        sheet: Worksheet, 
        cell_range: Optional[str] = None,
        line_type: str = "thin",
        line_color: str = "000000"
        ) -> None: 
        """
        Formatting borders from given range. 

        Args: 
            sheet (Worksheet): 
                Sheet required to perform given required actions
            cell_range (opt | str):
                Specificed start and end columns and rows (*ex. "A1:AA100"*)
            line_type (str):
                Line format to return, (*default is "thin"*)
            line_color (str):
                Line color to return, must specify the hex color 
        """
        # --- Line Type Check --- 
        valid_line_type = ["thin", "medium", "thick", "dashed", "dotted", "double", "hair"]
        if line_type not in valid_line_type: 
            line_type = "thin"
        # --- Styling the Line --- 
        line_type = Side(style=line_type, color=line_color)
        # Application to the cell --- 
        border = Border(left=line_type, right=line_type, top=line_type, bottom=line_type)
        # --- Range Arg --- 
        if cell_range: 
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
            for row in sheet.iter_rows(
                min_row=min_row, 
                max_row=max_row, 
                min_col=min_col, 
                max_col=max_col
            ):
                for cell in row: 
                    cell.border = border
        else:
            for row in sheet.iter_rows():
                for cell in row: 
                    cell.border = border

    def autofit_columns(
        self, 
        sheet: Worksheet,
        exceptions: Optional[set[str]] = None
    ) -> None: 
        """
        Resize all columns based on their contents. 

        Args: 
            sheet (Worksheet): 
                Sheet required to perform given required actions
            exceptions (opt | set[str]):
                Columns whose width is limited to the header length.
        """
        for column in sheet.columns: 
            column_letter = get_column_letter(column[0].column)
            header = column[0].value 
            
            # --- Exceptions Check --- 
            if exceptions and header in exceptions: 
                sheet.column_dimensions[column_letter].width = len(str(header)) + 2
                continue
                
            # --- Calculate Max Length of Contents ---
            max_length = max(
                (
                    len(str(cell.value)) 
                    for cell in column 
                    if cell.value is not None
                ), 
                default=0
            )
            sheet.column_dimensions[column_letter].width = max_length + 2